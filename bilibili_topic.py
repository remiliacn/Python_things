# Author: remiliacn
# This script should not be used for malicious reasons.
# Made for valletta's usage. (https://space.bilibili.com/573959614)

import asyncio
import csv
import json
import re
import time
from datetime import datetime, timedelta
from os import getcwd
from os.path import exists
from re import sub

import aiohttp
from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


# Decorator for calculating time usage.
def timer(func):
    async def wrapper(*args, **kwargs):
        start_time = time.time()
        await func(*args, **kwargs)
        end_time = time.time()
        logger.info(f'用时：{end_time - start_time:.2f}s')

    return wrapper


"""
BilibiliCrawl class

Takes a keyword (str), and a limit (int) parameter to start crawling bilibili topics
Reset is an optional parameter, which if is true, the crawling will start from scratch.
Uses asyncio to make crawling faster.
"""


class BilibiliCrawl:
    def __init__(self, keyword: str, limit: int, reset=False):
        self.keyword = keyword
        self.limit = limit
        self.reset = reset

        if not reset:
            # Only fetch one-day-old articles.
            self.time_limit = datetime.timestamp(
                datetime.now() - timedelta(days=1)
            )
        else:
            # Fetch 2-months-old work
            self.time_limit = datetime.timestamp(
                datetime.now() - timedelta(days=60)
            )

        self.total_page = None
        # This set is for checking if the article is already existed.
        self.id_list = set()

        self.api_link = f'https://api.bilibili.com/x/web-interface/search/' \
                        f'type?search_type=article&order=pubdate&keyword={self.keyword}'
        self.article_api_baseurl = 'https://www.bilibili.com/read/cv'
        self.csv_path = f'{getcwd()}/data/bot/valletta.csv'

        # Get time now for making the xlsx later.
        time_now = (datetime.utcnow() - timedelta(hours=8)).strftime("%Y-%m-%d_%H_%M_%S")
        self.xlsx_path = f'{getcwd()}/data/bot/valletta_{time_now}.xlsx'
        self.headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                          'AppleWebKit/537.36 (KHTML, like Gecko) '
                          'Chrome/91.0.4472.124 Safari/537.36'
        }

        self.new_articles = 0
        self.old_articles = 0
        self.word_not_enough_articles = 0

    def read_from_csv(self):
        # Read csv file if it exists, otherwise, create one.
        # If reset is true, ignore the file existing check.
        if not self.reset and exists(self.csv_path):
            with open(self.csv_path, encoding='utf_8_sig') as file:
                file_data_line = file.readlines()
                for x in file_data_line:
                    info = x.partition(',')[0]
                    if info.isdigit():
                        self.id_list.add(str(info))
        else:
            with open(self.csv_path, 'w+', encoding='utf_8_sig') as file:
                file.write(
                    ','.join(
                        [
                            '文章id',
                            '标题',
                            '发布时间',
                            '字数',
                            '文章url',
                            '评论数',
                            '点赞数',
                            '观看数',
                            '硬币数',
                            '分享数',
                            # '文章正文'
                        ]
                    ) + '\n'
                )

    async def _get_article_data(self, page=1) -> list:
        # Read API data from bilibili, and return articles for further analysis.
        async with aiohttp.ClientSession() as session:
            async with session.get(self.api_link + f'&page={page}', headers=self.headers) as response:
                data = await response.json()
                self.total_page = data['data']['numPages']
                logger.success(f'page {page} api fetch done.')
                data = data['data']['result']
                return [x for x in data if x['pub_time'] >= self.time_limit]

    async def _build_info(self):
        self.read_from_csv()
        await self._get_article_data()
        await asyncio.sleep(0.1)

        article_real_data = []
        for i in range(1, self.total_page + 1):
            article_data_single = await self._get_article_data(page=i)
            # Break loop if no article satisfies the time requirement.
            if not article_data_single:
                break

            if article_data_single is not None:
                for element in article_data_single:
                    article_id = str(element['id']).strip()
                    if article_id in self.id_list:
                        self.old_articles += 1
                        continue

                    article_real_data.append(element)

            await asyncio.sleep(0.1)

        tasks = []
        logger.debug('Starting individual page crawling threads...')
        for element in article_real_data:
            tasks.append(self.write_csv_worker(element))

        info = await asyncio.gather(*tasks)
        return info

    async def write_csv(self) -> (str, str):
        info = await self._build_info()

        logger.debug('Writing csv file...')
        with open(self.csv_path, 'a', encoding='utf-8-sig') as file:
            file.write(''.join([x for x in info if x]))

        logger.debug('Generating xlsx file with styles...')
        self._beautify_worksheet()
        return self.xlsx_path, await self.get_report()

    def _beautify_worksheet(self):
        with open(self.csv_path, encoding='utf_8_sig') as file:
            reader = csv.reader(file)
            wb = Workbook()
            ws = wb.active

            alignment_style = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for row_index, row in enumerate(reader):
                for column_index, cell in enumerate(row):
                    cell = ws.cell(column=column_index + 1, row=row_index + 1, value=cell)
                    cell.alignment = alignment_style

            # Applying table formats.
            ws.column_dimensions['B'].width = ws.column_dimensions['E'].width = 50
            ws.column_dimensions['C'].width = 20
            # ws.column_dimensions['K'].width = 40

            mediumStyle = TableStyleInfo(
                name='TableStyleMedium2',
                showRowStripes=True
            )

            total_article = self.new_articles + self.old_articles + self.word_not_enough_articles
            table = Table(
                displayName='Data',
                ref=f'A1:J{total_article + 1}',
                tableStyleInfo=mediumStyle
            )
            ws.add_table(table)
            wb.save(self.xlsx_path)

    async def get_report(self):
        return f'新加入文章：{self.new_articles}篇\n' \
               f'已加入文章：{self.old_articles}篇\n' \
               f'字数不够文章（已过滤）：{self.word_not_enough_articles}篇'

    async def write_csv_worker(self, data):
        article_id = str(data['id'])
        article_url = f'https://www.bilibili.com/read/cv{article_id}'

        async with aiohttp.ClientSession() as session:
            async with session.get(self.article_api_baseurl + article_id, headers=self.headers) as response:
                page_data = await response.text()
                page_data_info = re.findall(r'window.__INITIAL_STATE__=({.*?});', page_data)
                page_data_json = json.loads(page_data_info[0])

                word_count = page_data_json['readInfo']['words']

                if word_count < self.limit:
                    self.word_not_enough_articles += 1
                    return

                # word_list = [
                #     sub('(<.*?>|\xa0)', '', x)
                #     for x in re.findall(r'<p.*?>(.*?)</p>', page_data) if x
                # ]
                # article = '"' + "\n".join(word_list).strip() + '"'

                read_info_stat = page_data_json['readInfo']['stats']

                coin_count = read_info_stat['coin']
                share_count = read_info_stat['share']

        self.new_articles += 1
        title = sub(r'<.*?>', '', data['title']).replace('\n', '')
        pub_time = str(
            datetime.fromtimestamp(data['pub_time']).strftime('%Y-%m-%d %H:%M:%S')
        )
        reply = data['reply']
        like = data['like']
        view = data['view']

        logger.success(f'Task: {article_id} done.')

        return \
            ','.join(
                [
                    article_id, title,
                    pub_time, str(word_count),
                    article_url, str(reply),
                    str(like), str(view),
                    str(coin_count), str(share_count),
                    # article
                ]
            ) + '\n'


@timer
async def main():
    logger.debug('crawler starts')
    api = BilibiliCrawl('瓦莱塔学会', 800)
    path, report = await api.write_csv()
    logger.info(f'\n{report}', path)


if __name__ == '__main__':
    loop = asyncio.get_event_loop()
    loop.run_until_complete(main())
